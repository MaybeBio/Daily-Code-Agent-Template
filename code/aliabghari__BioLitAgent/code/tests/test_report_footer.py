"""
tests/test_report_footer.py
---------------------------
Unit tests verifying the Word digest report footer, section running footers,
and provenance retrieval from the knowledge store.
"""

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from docx import Document

from biolitagent.agents.report_generator import (
    _add_report_footer,
    _get_provenance_info,
)


class TestReportFooter(unittest.TestCase):
    def test_report_footer_provenance(self):
        temp_dir = tempfile.TemporaryDirectory()
        db_path = Path(temp_dir.name) / "test_report_db.db"
        conn = sqlite3.connect(str(db_path))
        conn.row_factory = sqlite3.Row
        conn.execute("CREATE TABLE knowledge_items (store_id INTEGER PRIMARY KEY, metadata_json TEXT)")
        conn.execute(
            "INSERT INTO knowledge_items (metadata_json) VALUES (?)",
            (json.dumps({
                "model_name": "gemini-2.5-flash",
                "config_versions": {"taxonomy_version": "1.1", "queries_version": "1.1"},
            }),)
        )
        conn.commit()

        try:
            model_name, tax_ver, queries_ver = _get_provenance_info(conn)
            self.assertEqual(model_name, "gemini-2.5-flash")
            self.assertEqual(tax_ver, "1.1")
            self.assertEqual(queries_ver, "1.1")

            doc = Document()
            _add_report_footer(doc, "2026-09-27T12:00:00Z", model_name, tax_ver, queries_ver)

            footer_text = doc.sections[0].footer.paragraphs[0].text
            self.assertIn("gemini-2.5-flash", footer_text)
            self.assertIn("v1.1", footer_text)

            body_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("BioLitAgent Digest Provenance", body_text)
            self.assertIn("Model: gemini-2.5-flash", body_text)
            conn.close()
        finally:
            temp_dir.cleanup()

    def test_report_footer_fallback_when_empty_db(self):
        temp_dir = tempfile.TemporaryDirectory()
        db_path = Path(temp_dir.name) / "empty_db.db"
        conn = sqlite3.connect(str(db_path))
        conn.row_factory = sqlite3.Row
        conn.execute("CREATE TABLE knowledge_items (store_id INTEGER PRIMARY KEY, metadata_json TEXT)")
        conn.commit()

        try:
            model_name, tax_ver, queries_ver = _get_provenance_info(conn)
            self.assertEqual(model_name, "gemini-2.5-flash")
            self.assertEqual(tax_ver, "1.1")
            self.assertEqual(queries_ver, "1.1")

            doc = Document()
            _add_report_footer(doc, "2026-09-27T12:00:00Z", model_name, tax_ver, queries_ver)
            footer_text = doc.sections[0].footer.paragraphs[0].text
            self.assertIn("Model: gemini-2.5-flash", footer_text)
            conn.close()
        finally:
            temp_dir.cleanup()


if __name__ == "__main__":
    unittest.main()
